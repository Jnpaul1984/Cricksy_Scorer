"""Secure staged CSV/XLSX imports into the School master roster."""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import re
import unicodedata
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import structlog
from backend.api.schemas.player_imports import (
    PlayerImportApplyRequest,
    PlayerImportApplyResponse,
    PlayerImportApplySummary,
    PlayerImportPreviewResponse,
)
from backend.services import organization_service
from backend.services.organization_entitlement_service import (
    require_organization_capability,
)
from backend.sql_app.models import (
    PlayerProfile,
    SchoolPlayerImport,
    SchoolPlayerMembership,
    SchoolTeamPlayerMembership,
    Team,
)
from fastapi import UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

logger = structlog.get_logger(__name__)

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 32
MAX_CELL_LENGTH = 255
MAX_XLSX_ARCHIVE_ENTRIES = 512
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 200
PREVIEW_TTL = dt.timedelta(minutes=30)

ROSTER_CAPABILITY = "school_master_roster"
TEAM_CAPABILITIES = ("school_persistent_teams", "school_team_rosters")
IMPORT_ROLES = frozenset({"owner", "admin", "coach"})
LIFECYCLE_ROLES = frozenset({"owner", "admin"})
ALLOWED_FIELDS = frozenset({"player_name", "student_identifier", "year_group", "team_name"})
DIRECT_HEADERS = {
    "player_name": "player_name",
    "student_identifier": "student_identifier",
    "year_group": "year_group",
    "team": "team_name",
    "team_name": "team_name",
}


@dataclass(frozen=True)
class SchoolPlayerImportServiceError(Exception):
    status_code: int
    detail: str


class _RowApplyError(Exception):
    pass


def _error(status_code: int, detail: str) -> SchoolPlayerImportServiceError:
    return SchoolPlayerImportServiceError(status_code, detail)


def _new_id() -> str:
    return str(uuid.uuid4())


def _normalized_header(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).strip().split()).casefold()


def _normalized_name(value: str) -> str:
    return " ".join(value.split()).casefold()


def _safe_filename(value: str | None) -> str:
    basename = Path(value or "").name
    normalized = unicodedata.normalize("NFKC", basename)
    sanitized = re.sub(r"[^A-Za-z0-9._-]", "_", normalized).strip("._")
    return (sanitized or "roster")[:255]


def _clean_value(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    if not text:
        return None
    if len(text) > MAX_CELL_LENGTH:
        raise _error(422, f"Cell value exceeds {MAX_CELL_LENGTH} characters")
    return text


def _reject_formula_text(value: str | None) -> None:
    if value is not None and value.lstrip().startswith(("=", "+", "-", "@")):
        raise _error(422, "Formula-like cell values are not supported")


def _validate_archive(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
                raise _error(422, "XLSX archive contains too many entries")
            total_size = 0
            for entry in entries:
                normalized_name = entry.filename.replace("\\", "/").casefold()
                if entry.flag_bits & 0x1:
                    raise _error(422, "Encrypted XLSX files are not supported")
                if normalized_name.endswith("vbaProject.bin".casefold()):
                    raise _error(422, "Macro-enabled workbooks are not supported")
                if normalized_name.startswith("xl/externallinks/"):
                    raise _error(422, "External workbook links are not supported")
                total_size += entry.file_size
                if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise _error(422, "XLSX expanded size exceeds the safety limit")
                if (
                    entry.file_size > 1_000_000
                    and entry.compress_size > 0
                    and entry.file_size / entry.compress_size > MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise _error(422, "XLSX compression ratio exceeds the safety limit")
            if archive.testzip() is not None:
                raise _error(422, "Malformed XLSX archive")
    except SchoolPlayerImportServiceError:
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:
        raise _error(422, "Malformed or encrypted XLSX file") from exc


def _parse_csv(payload: bytes) -> list[list[str | None]]:
    if b"\x00" in payload:
        raise _error(422, "Binary content is not valid CSV")
    try:
        text = payload.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise _error(422, "CSV must use UTF-8 encoding") from exc
    try:
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        rows: list[list[str | None]] = []
        for raw_row in reader:
            if len(raw_row) > MAX_COLUMNS:
                raise _error(422, f"Import supports at most {MAX_COLUMNS} columns")
            row = [_clean_value(value) for value in raw_row]
            for value in row:
                _reject_formula_text(value)
            if any(value is not None for value in row):
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise _error(422, f"Import supports at most {MAX_ROWS} data rows")
    except csv.Error as exc:
        raise _error(422, "Malformed CSV file") from exc
    return rows


def _parse_xlsx(payload: bytes) -> list[list[str | None]]:
    _validate_archive(payload)
    try:
        workbook = load_workbook(
            io.BytesIO(payload),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (InvalidFileException, KeyError, OSError, ValueError, RuntimeError) as exc:
        raise _error(422, "Malformed or encrypted XLSX file") from exc
    try:
        if len(workbook.worksheets) != 1:
            raise _error(422, "XLSX imports must contain exactly one worksheet")
        worksheet = workbook.worksheets[0]
        if worksheet.max_column > MAX_COLUMNS:
            raise _error(422, f"Import supports at most {MAX_COLUMNS} columns")
        if worksheet.max_row > MAX_ROWS + 1:
            raise _error(422, f"Import supports at most {MAX_ROWS} data rows")
        rows: list[list[str | None]] = []
        for raw_row in worksheet.iter_rows():
            row: list[str | None] = []
            for cell in raw_row:
                if cell.data_type == "f":
                    raise _error(422, "XLSX formulas are not supported")
                value = _clean_value(cell.value)
                _reject_formula_text(value)
                row.append(value)
            if any(value is not None for value in row):
                rows.append(row)
        return rows
    finally:
        workbook.close()


def _parse_mapping(mapping_json: str | None) -> dict[str, str]:
    if mapping_json is None or not mapping_json.strip():
        return {}
    try:
        parsed = json.loads(mapping_json)
    except json.JSONDecodeError as exc:
        raise _error(422, "column_mapping must be valid JSON") from exc
    if not isinstance(parsed, dict):
        raise _error(422, "column_mapping must be a JSON object")
    mapping: dict[str, str] = {}
    for source, target in parsed.items():
        if not isinstance(source, str) or not isinstance(target, str):
            raise _error(422, "column_mapping keys and values must be strings")
        canonical = "team_name" if target == "team" else target
        if canonical not in ALLOWED_FIELDS:
            raise _error(422, f"Unsupported mapped field: {target}")
        normalized_source = _normalized_header(source)
        if not normalized_source:
            raise _error(422, "Mapped source headers must not be empty")
        mapping[normalized_source] = canonical
    return mapping


def _map_rows(
    rows: list[list[str | None]],
    mapping_json: str | None,
) -> tuple[dict[str, str], list[tuple[int, dict[str, str | None]]]]:
    if not rows:
        raise _error(422, "Import file is empty")
    headers = rows[0]
    if not headers or all(value is None for value in headers):
        raise _error(422, "Import header row is empty")
    header_names = [value or "" for value in headers]
    normalized_headers = [_normalized_header(value) for value in header_names]
    if any(not value for value in normalized_headers):
        raise _error(422, "Import headers must not be empty")
    duplicates = [name for name, count in Counter(normalized_headers).items() if count > 1]
    if duplicates:
        raise _error(422, "Duplicate or ambiguous import headers are not supported")

    explicit = _parse_mapping(mapping_json)
    unknown_sources = sorted(set(explicit) - set(normalized_headers))
    if unknown_sources:
        raise _error(422, "column_mapping references a header not present in the file")

    resolved: dict[str, str] = {}
    targets: set[str] = set()
    for original, normalized in zip(header_names, normalized_headers, strict=True):
        target = explicit.get(normalized, DIRECT_HEADERS.get(normalized))
        if target is None:
            continue
        if target in targets:
            raise _error(422, f"Multiple columns map to {target}")
        targets.add(target)
        resolved[original] = target
    if "player_name" not in targets:
        raise _error(422, "A column must map to required field player_name")

    mapped_rows: list[tuple[int, dict[str, str | None]]] = []
    for source_row_number, row in enumerate(rows[1:], start=2):
        padded = row + [None] * (len(headers) - len(row))
        values = dict.fromkeys(ALLOWED_FIELDS)
        for index, original in enumerate(header_names):
            target = resolved.get(original)
            if target is not None:
                values[target] = padded[index] if index < len(padded) else None
        mapped_rows.append((source_row_number, values))
    if not mapped_rows:
        raise _error(422, "Import contains no player rows")
    return resolved, mapped_rows


async def _require_import_role(
    db: AsyncSession,
    *,
    organization_id: str,
    actor_user_id: str,
    require_team_capabilities: bool = False,
) -> str:
    await require_organization_capability(
        db,
        organization_id=organization_id,
        actor_user_id=actor_user_id,
        capability=ROSTER_CAPABILITY,
    )
    _, membership = await organization_service.get_organization_for_member(
        db,
        organization_id=organization_id,
        user_id=actor_user_id,
    )
    if membership.role not in IMPORT_ROLES:
        raise _error(403, "Insufficient organization role")
    if require_team_capabilities:
        for capability in TEAM_CAPABILITIES:
            await require_organization_capability(
                db,
                organization_id=organization_id,
                actor_user_id=actor_user_id,
                capability=capability,
            )
    return membership.role


async def _build_preview_rows(
    db: AsyncSession,
    *,
    organization_id: str,
    mapped_rows: list[tuple[int, dict[str, str | None]]],
) -> list[dict[str, Any]]:
    roster_result = await db.execute(
        select(SchoolPlayerMembership, PlayerProfile)
        .join(PlayerProfile, PlayerProfile.player_id == SchoolPlayerMembership.player_profile_id)
        .where(SchoolPlayerMembership.organization_id == organization_id)
    )
    roster = list(roster_result.tuples().all())
    team_result = await db.execute(select(Team).where(Team.organization_id == organization_id))
    teams = list(team_result.scalars().all())

    identifiers = Counter(
        values["student_identifier"] for _, values in mapped_rows if values["student_identifier"]
    )
    names = Counter(
        _normalized_name(values["player_name"] or "")
        for _, values in mapped_rows
        if values["player_name"]
    )
    preview_rows: list[dict[str, Any]] = []

    for source_row_number, values in mapped_rows:
        errors: list[str] = []
        warnings: list[str] = []
        ambiguity_reason: str | None = None
        player_name = values["player_name"]
        student_identifier = values["student_identifier"]
        team_name = values["team_name"]
        if not player_name:
            errors.append("player_name is required")

        candidate_rows: list[tuple[SchoolPlayerMembership, PlayerProfile]] = []
        if student_identifier:
            candidate_rows.extend(
                row for row in roster if row[0].student_identifier == student_identifier
            )
            if identifiers[student_identifier] > 1:
                ambiguity_reason = "student_identifier appears more than once in the file"
        if player_name:
            same_name = [
                row
                for row in roster
                if _normalized_name(row[1].player_name) == _normalized_name(player_name)
            ]
            for row in same_name:
                if row not in candidate_rows:
                    candidate_rows.append(row)
            if names[_normalized_name(player_name)] > 1:
                ambiguity_reason = "same player name appears more than once in the file"

        candidate_memberships = [
            {
                "school_player_membership_id": membership.id,
                "player_name": profile.player_name,
                "status": membership.status,
            }
            for membership, profile in candidate_rows
        ]
        exact_identifier_candidates = (
            [row for row in candidate_rows if row[0].student_identifier == student_identifier]
            if student_identifier
            else []
        )

        team_candidates: list[dict[str, str]] = []
        resolved_team_id: str | None = None
        if team_name:
            matching_teams = [
                team for team in teams if _normalized_name(team.name) == _normalized_name(team_name)
            ]
            active_teams = [team for team in matching_teams if team.status == "active"]
            team_candidates = [
                {"team_id": team.id, "team_name": team.name} for team in active_teams
            ]
            if len(active_teams) == 1:
                resolved_team_id = active_teams[0].id
            elif len(active_teams) > 1:
                ambiguity_reason = "duplicate active Team names require exact Team selection"
            elif matching_teams:
                errors.append("Archived Team cannot receive roster assignments")
            else:
                errors.append("No active Team matches team_name")

        resolved_membership_id: str | None = None
        classification = "create_new"
        resolution_required = False
        if errors:
            classification = "invalid"
        elif len(exact_identifier_candidates) == 1 and player_name:
            membership, profile = exact_identifier_candidates[0]
            if _normalized_name(profile.player_name) != _normalized_name(player_name):
                classification = "ambiguous_needs_review"
                ambiguity_reason = "student_identifier candidate has a conflicting player name"
                resolution_required = True
            else:
                resolved_membership_id = membership.id
                resolution_required = True
                if membership.status == "inactive":
                    classification = "ambiguous_needs_review"
                    ambiguity_reason = "inactive membership requires explicit reactivation"
                elif team_name:
                    classification = "team_assignment_only"
                    ambiguity_reason = (
                        "student_identifier is a candidate signal requiring confirmation"
                    )
                else:
                    classification = "duplicate_existing_school_membership"
                    ambiguity_reason = (
                        "student_identifier is a candidate signal requiring confirmation"
                    )
        elif candidate_rows or ambiguity_reason:
            classification = "ambiguous_needs_review"
            resolution_required = True
            ambiguity_reason = (
                ambiguity_reason or "same-name or identifier candidates require review"
            )

        if team_name and resolved_team_id is None and not errors:
            classification = "ambiguous_needs_review"
            resolution_required = True
        if len(exact_identifier_candidates) > 1:
            warnings.append("student_identifier is not unique within this School")
            classification = "ambiguous_needs_review"
            resolution_required = True
            ambiguity_reason = "multiple School roster memberships share student_identifier"
        if candidate_rows and not student_identifier:
            warnings.append("Names are never treated as proof of identity")

        preview_rows.append(
            {
                "source_row_number": source_row_number,
                "values": values,
                "classification": classification,
                "validation_errors": errors,
                "warnings": warnings,
                "ambiguity_reason": ambiguity_reason,
                "resolution_required": resolution_required,
                "candidate_memberships": candidate_memberships,
                "resolved_school_player_membership_id": resolved_membership_id,
                "team_candidates": team_candidates,
                "resolved_team_id": resolved_team_id,
            }
        )
    return preview_rows


async def create_preview(
    db: AsyncSession,
    *,
    organization_id: str,
    actor_user_id: str,
    upload: UploadFile,
    column_mapping_json: str | None,
) -> PlayerImportPreviewResponse:
    await _require_import_role(
        db,
        organization_id=organization_id,
        actor_user_id=actor_user_id,
    )
    filename = _safe_filename(upload.filename)
    extension = Path(filename).suffix.casefold()
    if extension not in {".csv", ".xlsx"}:
        logger.info(
            "organization.player_import_preview_rejected",
            organization_id=organization_id,
            actor_user_id=actor_user_id,
            reason="unsupported_extension",
        )
        raise _error(415, "Only .csv and .xlsx files are supported")
    payload = await upload.read(MAX_FILE_BYTES + 1)
    await upload.close()
    if not payload:
        raise _error(422, "Import file is empty")
    if len(payload) > MAX_FILE_BYTES:
        raise _error(413, f"Import file exceeds {MAX_FILE_BYTES} bytes")

    file_type = extension[1:]
    try:
        raw_rows = _parse_csv(payload) if file_type == "csv" else _parse_xlsx(payload)
        column_mapping, mapped_rows = _map_rows(raw_rows, column_mapping_json)
        has_team_assignments = any(values["team_name"] for _, values in mapped_rows)
        if has_team_assignments:
            await _require_import_role(
                db,
                organization_id=organization_id,
                actor_user_id=actor_user_id,
                require_team_capabilities=True,
            )
        preview_rows = await _build_preview_rows(
            db,
            organization_id=organization_id,
            mapped_rows=mapped_rows,
        )
    except Exception:
        logger.info(
            "organization.player_import_preview_rejected",
            organization_id=organization_id,
            actor_user_id=actor_user_id,
            file_type=file_type,
        )
        raise

    now = dt.datetime.now(dt.UTC)
    import_session = SchoolPlayerImport(
        id=_new_id(),
        organization_id=organization_id,
        actor_user_id=actor_user_id,
        status="previewed",
        file_type=file_type,
        original_filename=filename or f"roster.{file_type}",
        content_sha256=hashlib.sha256(payload).hexdigest(),
        column_mapping=column_mapping,
        preview_rows=preview_rows,
        row_count=len(preview_rows),
        expires_at=now + PREVIEW_TTL,
    )
    db.add(import_session)
    await db.commit()
    await db.refresh(import_session)
    logger.info(
        "organization.player_import_preview_created",
        organization_id=organization_id,
        actor_user_id=actor_user_id,
        import_id=import_session.id,
        file_type=file_type,
        row_count=len(preview_rows),
    )
    return PlayerImportPreviewResponse(
        import_id=import_session.id,
        file_type=file_type,
        original_filename=import_session.original_filename,
        content_sha256=import_session.content_sha256,
        row_count=import_session.row_count,
        column_mapping=import_session.column_mapping,
        expires_at=import_session.expires_at,
        rows=import_session.preview_rows,
    )


async def _load_membership(
    db: AsyncSession,
    *,
    organization_id: str,
    membership_id: str,
) -> SchoolPlayerMembership:
    result = await db.execute(
        select(SchoolPlayerMembership).where(
            SchoolPlayerMembership.id == membership_id,
            SchoolPlayerMembership.organization_id == organization_id,
        )
    )
    membership = result.scalar_one_or_none()
    if membership is None:
        raise _RowApplyError("Roster player is not available in this organization")
    return membership


async def _load_team(
    db: AsyncSession,
    *,
    organization_id: str,
    team_id: str,
) -> Team:
    result = await db.execute(
        select(Team).where(Team.id == team_id, Team.organization_id == organization_id)
    )
    team = result.scalar_one_or_none()
    if team is None:
        raise _RowApplyError("Team is not available in this organization")
    if team.status != "active":
        raise _RowApplyError("Archived Team cannot receive roster assignments")
    return team


async def _assign_team(
    db: AsyncSession,
    *,
    organization_id: str,
    team_id: str,
    school_player_membership_id: str,
    actor_user_id: str,
) -> tuple[str, SchoolTeamPlayerMembership]:
    await _load_team(db, organization_id=organization_id, team_id=team_id)
    existing_result = await db.execute(
        select(SchoolTeamPlayerMembership).where(
            SchoolTeamPlayerMembership.organization_id == organization_id,
            SchoolTeamPlayerMembership.team_id == team_id,
            SchoolTeamPlayerMembership.school_player_membership_id == school_player_membership_id,
        )
    )
    existing = existing_result.scalar_one_or_none()
    if existing is not None:
        if existing.status == "inactive":
            existing.status = "active"
            return "reactivated", existing
        return "no_op", existing
    assignment = SchoolTeamPlayerMembership(
        id=_new_id(),
        organization_id=organization_id,
        team_id=team_id,
        school_player_membership_id=school_player_membership_id,
        status="active",
        created_by_user_id=actor_user_id,
    )
    db.add(assignment)
    return "created", assignment


def _allowed_candidate_ids(row: dict[str, Any]) -> set[str]:
    return {
        str(candidate["school_player_membership_id"]) for candidate in row["candidate_memberships"]
    }


def _allowed_team_ids(row: dict[str, Any]) -> set[str]:
    allowed = {str(candidate["team_id"]) for candidate in row["team_candidates"]}
    if row.get("resolved_team_id"):
        allowed.add(str(row["resolved_team_id"]))
    return allowed


async def apply_import(
    db: AsyncSession,
    *,
    organization_id: str,
    import_id: str,
    actor_user_id: str,
    payload: PlayerImportApplyRequest,
) -> PlayerImportApplyResponse:
    role = await _require_import_role(
        db,
        organization_id=organization_id,
        actor_user_id=actor_user_id,
    )
    result = await db.execute(
        select(SchoolPlayerImport)
        .where(
            SchoolPlayerImport.id == import_id,
            SchoolPlayerImport.organization_id == organization_id,
            SchoolPlayerImport.actor_user_id == actor_user_id,
        )
        .with_for_update()
    )
    import_session = result.scalar_one_or_none()
    if import_session is None:
        logger.warning(
            "organization.player_import_tenant_violation",
            organization_id=organization_id,
            actor_user_id=actor_user_id,
            import_id=import_id,
        )
        raise _error(404, "Player import not found")
    if import_session.status == "applied":
        logger.info(
            "organization.player_import_apply_conflict",
            organization_id=organization_id,
            actor_user_id=actor_user_id,
            import_id=import_id,
        )
        raise _error(409, "Player import has already been applied")
    expires_at = import_session.expires_at
    if expires_at.tzinfo is None:
        expires_at = expires_at.replace(tzinfo=dt.UTC)
    if expires_at <= dt.datetime.now(dt.UTC):
        raise _error(410, "Player import preview has expired")

    preview_by_row = {row["source_row_number"]: row for row in import_session.preview_rows}
    resolution_by_row = {item.source_row_number: item for item in payload.resolutions}
    if set(resolution_by_row) - set(preview_by_row):
        raise _error(422, "Resolution references a row outside this preview")
    has_team_assignments = any(
        row["values"].get("team_name") for row in import_session.preview_rows
    )
    if has_team_assignments:
        await _require_import_role(
            db,
            organization_id=organization_id,
            actor_user_id=actor_user_id,
            require_team_capabilities=True,
        )

    logger.info(
        "organization.player_import_apply_started",
        organization_id=organization_id,
        actor_user_id=actor_user_id,
        import_id=import_id,
        row_count=import_session.row_count,
    )
    summary = {
        "created_players": 0,
        "linked_existing_players": 0,
        "reactivated_memberships": 0,
        "team_assignments_created": 0,
        "team_assignments_reactivated": 0,
        "no_op_rows": 0,
        "skipped_rows": 0,
        "failed_rows": 0,
    }
    row_results: list[dict[str, Any]] = []

    for source_row_number in sorted(preview_by_row):
        row = preview_by_row[source_row_number]
        resolution = resolution_by_row.get(source_row_number)
        if row["classification"] == "invalid":
            summary["failed_rows"] += 1
            row_results.append(
                {
                    "source_row_number": source_row_number,
                    "outcome": "failed",
                    "detail": "; ".join(row["validation_errors"]),
                }
            )
            continue
        if row["resolution_required"] and resolution is None:
            summary["failed_rows"] += 1
            row_results.append(
                {
                    "source_row_number": source_row_number,
                    "outcome": "failed",
                    "detail": "Explicit resolution is required",
                }
            )
            continue
        action = resolution.action if resolution is not None else "create_new"
        if action == "skip":
            summary["skipped_rows"] += 1
            row_results.append(
                {
                    "source_row_number": source_row_number,
                    "outcome": "skipped",
                    "detail": "Skipped by explicit resolution",
                }
            )
            continue

        values = row["values"]
        membership: SchoolPlayerMembership | None = None
        player_profile_id: str | None = None
        target_team_id: str | None = None
        try:
            async with db.begin_nested():
                if action == "create_new":
                    if row["resolution_required"] and resolution is None:
                        raise _RowApplyError("Explicit create_new resolution is required")
                    profile = PlayerProfile(
                        player_id=_new_id(),
                        player_name=values["player_name"],
                    )
                    membership = SchoolPlayerMembership(
                        id=_new_id(),
                        organization_id=organization_id,
                        player_profile_id=profile.player_id,
                        status="active",
                        student_identifier=values.get("student_identifier"),
                        year_group=values.get("year_group"),
                        created_by_user_id=actor_user_id,
                    )
                    db.add_all([profile, membership])
                    player_profile_id = profile.player_id
                    base_outcome = "created"
                elif action in {"use_existing", "reactivate_existing"}:
                    assert resolution is not None
                    membership_id = resolution.school_player_membership_id
                    if membership_id not in _allowed_candidate_ids(row):
                        raise _RowApplyError("Selected roster player was not part of this preview")
                    membership = await _load_membership(
                        db,
                        organization_id=organization_id,
                        membership_id=membership_id,
                    )
                    player_profile_id = membership.player_profile_id
                    if action == "reactivate_existing":
                        if role not in LIFECYCLE_ROLES:
                            raise _RowApplyError(
                                "Only owner/admin may reactivate roster membership"
                            )
                        if membership.status != "inactive":
                            raise _RowApplyError("Selected roster membership is not inactive")
                        membership.status = "active"
                        base_outcome = "reactivated"
                    else:
                        if membership.status != "active":
                            raise _RowApplyError(
                                "Inactive membership requires explicit reactivation"
                            )
                        base_outcome = "linked_existing"
                else:
                    raise _RowApplyError("Unsupported row resolution")

                if values.get("team_name"):
                    target_team_id = (
                        resolution.team_id
                        if resolution is not None
                        else row.get("resolved_team_id")
                    )
                    if not target_team_id:
                        raise _RowApplyError("Exact Team resolution is required")
                    if target_team_id not in _allowed_team_ids(row):
                        raise _RowApplyError("Selected Team was not part of this preview")
                    assert membership is not None
                    team_action, _ = await _assign_team(
                        db,
                        organization_id=organization_id,
                        team_id=target_team_id,
                        school_player_membership_id=membership.id,
                        actor_user_id=actor_user_id,
                    )
                else:
                    if resolution is not None and resolution.team_id is not None:
                        raise _RowApplyError("Team selection was not part of this preview")
                    team_action = None
                await db.flush()

            if base_outcome == "created":
                summary["created_players"] += 1
            elif base_outcome == "reactivated":
                summary["reactivated_memberships"] += 1
            else:
                summary["linked_existing_players"] += 1
            if team_action == "created":
                summary["team_assignments_created"] += 1
                outcome = "team_assigned"
            elif team_action == "reactivated":
                summary["team_assignments_reactivated"] += 1
                outcome = "team_assigned"
            elif team_action == "no_op":
                summary["no_op_rows"] += 1
                outcome = "no_op"
            else:
                outcome = base_outcome
            row_results.append(
                {
                    "source_row_number": source_row_number,
                    "outcome": outcome,
                    "school_player_membership_id": membership.id if membership else None,
                    "player_profile_id": player_profile_id,
                    "team_id": target_team_id,
                    "detail": "Row applied successfully",
                }
            )
        except (_RowApplyError, IntegrityError) as exc:
            summary["failed_rows"] += 1
            logger.info(
                "organization.player_import_row_failed",
                organization_id=organization_id,
                actor_user_id=actor_user_id,
                import_id=import_id,
                source_row_number=source_row_number,
                reason=type(exc).__name__,
            )
            row_results.append(
                {
                    "source_row_number": source_row_number,
                    "outcome": "failed",
                    "detail": str(exc) if isinstance(exc, _RowApplyError) else "Database conflict",
                }
            )

    applied_at = dt.datetime.now(dt.UTC)
    stored_result = {
        "import_id": import_id,
        "status": "applied",
        "applied_at": applied_at.isoformat(),
        "summary": summary,
        "rows": row_results,
    }
    import_session.status = "applied"
    import_session.applied_at = applied_at
    import_session.result = stored_result
    await db.commit()
    logger.info(
        "organization.player_import_apply_completed",
        organization_id=organization_id,
        actor_user_id=actor_user_id,
        import_id=import_id,
        **summary,
    )
    return PlayerImportApplyResponse(
        import_id=import_id,
        status="applied",
        applied_at=applied_at,
        summary=PlayerImportApplySummary(**summary),
        rows=row_results,
    )
